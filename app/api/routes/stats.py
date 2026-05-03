from datetime import datetime
from io import BytesIO
from typing import Optional

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Border, Font, PatternFill, Side
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.db.session import get_session
from app.models.order import ORDER_STATUSES, Order, OrderItem
from app.models.product import Product
from app.models.product_stats import ProductStats
from app.models.user import User

router = APIRouter(prefix="/stats", tags=["stats"])


async def ensure_stats(product_id: int, session: AsyncSession) -> ProductStats:
    res = await session.execute(
        select(ProductStats).where(ProductStats.product_id == product_id)
    )
    s = res.scalar_one_or_none()
    if not s:
        s = ProductStats(product_id=product_id)
        session.add(s)
        await session.flush()
    return s


@router.get("/dashboard")
async def get_dashboard(
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    session: AsyncSession = Depends(get_session),
):

    orders_q = select(Order).order_by(Order.created_at.desc())
    if date_from:
        orders_q = orders_q.where(Order.created_at >= datetime.fromisoformat(date_from))
    if date_to:
        orders_q = orders_q.where(Order.created_at <= datetime.fromisoformat(date_to + "T23:59:59"))

    orders_res = await session.execute(orders_q)
    all_orders = orders_res.scalars().all()

    billable_orders = [o for o in all_orders if o.status not in ("cancelled", "returned")]
    total_revenue = sum(o.total for o in billable_orders)
    total_orders = len(all_orders)
    billable_orders_count = len(billable_orders)
    average_order_value = int(total_revenue / billable_orders_count) if billable_orders_count else 0
    by_status = {}
    for o in all_orders:
        by_status[o.status] = by_status.get(o.status, 0) + 1

    user_map = {}
    user_ids = {o.user_id for o in all_orders}
    if user_ids:
        users_res = await session.execute(select(User).where(User.id.in_(user_ids)))
        users = users_res.scalars().all()
        user_map = {u.id: u for u in users}

    status_order = list(ORDER_STATUSES.keys())
    sorted_statuses = [s for s in status_order if s in by_status] + [s for s in by_status if s not in status_order]

    recent_orders = []
    for order in all_orders[:10]:
        user = user_map.get(order.user_id)
        user_name = f"ID:{order.user_id}"
        if user:
            user_name = f"{user.first_name or ''} {user.last_name or ''}".strip() or user.username or user_name

        recent_orders.append({
            "id": order.id,
            "user_id": order.user_id,
            "user_name": user_name,
            "status": order.status,
            "status_label": ORDER_STATUSES.get(order.status, order.status),
            "total": order.total,
            "comment": order.comment,
            "delivery_address": order.delivery_address,
            "created_at": order.created_at.isoformat() if order.created_at else None,
        })

    return {
        "period": {"from": date_from, "to": date_to},
        "total_revenue": total_revenue,
        "total_orders": total_orders,
        "billable_orders": billable_orders_count,
        "average_order_value": average_order_value,
        "orders_by_status": [
            {"status": status, "label": ORDER_STATUSES.get(status, status), "count": by_status[status]}
            for status in sorted_statuses
        ],
        "recent_orders": recent_orders,
    }


@router.get("/products")
async def get_all_stats(
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    session: AsyncSession = Depends(get_session),
):
    dashboard_orders = 0
    try:
        orders_q = select(func.count(Order.id)).where(~Order.status.in_(["cancelled", "returned"]))
        if date_from:
            orders_q = orders_q.where(Order.created_at >= datetime.fromisoformat(date_from))
        if date_to:
            orders_q = orders_q.where(Order.created_at <= datetime.fromisoformat(date_to + "T23:59:59"))
        dashboard_orders = (await session.execute(orders_q)).scalar() or 0
    except Exception:
        pass

    res = await session.execute(select(ProductStats))
    stats = {s.product_id: s for s in res.scalars().all()}

    query = select(
        OrderItem.product_id,
        func.sum(OrderItem.quantity).label("sold_qty"),
        func.sum(OrderItem.price * OrderItem.quantity).label("sold_sum"),
    ).join(Order, OrderItem.order_id == Order.id).where(
        Order.status.notin_(["cancelled", "returned"])
    ).group_by(OrderItem.product_id)

    if date_from:
        query = query.where(Order.created_at >= datetime.fromisoformat(date_from))
    if date_to:
        query = query.where(Order.created_at <= datetime.fromisoformat(date_to + "T23:59:59"))

    period_res = await session.execute(query)
    period = {row.product_id: {"sold_qty": row.sold_qty, "sold_sum": row.sold_sum}
              for row in period_res.all()}

    result = []
    total_sold_qty = 0
    total_sold_sum = 0
    total_returned = 0
    all_pids = set(stats) | set(period)
    for pid in all_pids:
        s = stats.get(pid)
        p = period.get(pid, {})
        sold_qty = p.get("sold_qty", 0)
        sold_sum = p.get("sold_sum", 0)
        returned = s.returned if s else 0
        total_sold_qty += sold_qty
        total_sold_sum += sold_sum
        total_returned += returned
        result.append({
            "product_id": pid,
            "added_to_cart": s.added_to_cart if s else 0,
            "ordered": s.ordered if s else 0,
            "returned": returned,
            "period_sold_qty": sold_qty,
            "period_sold_sum": sold_sum,
        })

    products_with_sales = sum(1 for r in result if r["period_sold_qty"] > 0)
    avg_items_per_order = int(total_sold_qty / dashboard_orders) if dashboard_orders else 0
    avg_price = int(total_sold_sum / total_sold_qty) if total_sold_qty else 0

    return {
        "items": result,
        "summary": {
            "total_sold_sum": total_sold_sum,
            "total_sold_qty": total_sold_qty,
            "avg_items_per_order": avg_items_per_order,
            "avg_price": avg_price,
            "total_returned": total_returned,
            "products_with_sales": products_with_sales,
        }
    }


@router.post("/products/{product_id}/return")
async def track_return(product_id: int, session: AsyncSession = Depends(get_session)):
    s = await ensure_stats(product_id, session)
    s.returned += 1
    await session.commit()
    return {"ok": True}

@router.get("/dashboard/export")
async def export_dashboard(
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    session: AsyncSession = Depends(get_session),
):
    dashboard = await get_dashboard(date_from, date_to, session)

    wb = Workbook()
    ws = wb.active
    ws.title = "Статистика заказов"

    # Стили
    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill(start_color="6C63FF", end_color="6C63FF", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin'),
    )

    ws.merge_cells('A1:C1')
    ws['A1'] = f"Статистика заказов • Период: {date_from or '—'} – {date_to or '—'}"
    ws['A1'].font = Font(bold=True, size=14)

    ws['A3'] = "Показатель"; ws['B3'] = "Значение"
    for col in (1, 2):
        c = ws.cell(row=3, column=col)
        c.font = hdr_font; c.fill = hdr_fill; c.border = thin_border

    summary = [
        ("Выручка, ₽", dashboard['total_revenue']),
        ("Всего заказов", dashboard['total_orders']),
        ("Без отмен/возвратов", dashboard['billable_orders']),
        ("Средний чек, ₽", dashboard['average_order_value']),
    ]
    for i, (label, value) in enumerate(summary):
        ws.cell(row=4 + i, column=1, value=label).border = thin_border
        ws.cell(row=4 + i, column=2, value=value).border = thin_border

    ws['A10'] = "Статус"; ws['B10'] = "Количество"
    for col in (1, 2):
        c = ws.cell(row=10, column=col)
        c.font = hdr_font; c.fill = hdr_fill; c.border = thin_border

    for i, st in enumerate(dashboard.get('orders_by_status', [])):
        ws.cell(row=11 + i, column=1, value=st['label']).border = thin_border
        ws.cell(row=11 + i, column=2, value=st['count']).border = thin_border

    start = 11 + len(dashboard.get('orders_by_status', [])) + 2
    ws.cell(row=start, column=1, value="Последние заказы").font = Font(bold=True, size=12)
    start += 1
    headers = ["Заказ", "Покупатель", "Статус", "Сумма, ₽", "Дата"]
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=start, column=col, value=h)
        c.font = hdr_font; c.fill = hdr_fill; c.border = thin_border

    for i, order in enumerate(dashboard.get('recent_orders', [])):
        row = start + 1 + i
        ws.cell(row=row, column=1, value=f"#{order['id']}").border = thin_border
        ws.cell(row=row, column=2, value=order.get('user_name', '')).border = thin_border
        ws.cell(row=row, column=3, value=order.get('status_label', '')).border = thin_border
        ws.cell(row=row, column=4, value=order.get('total', 0)).border = thin_border
        ws.cell(row=row, column=4).number_format = '# ##0'
        ws.cell(row=row, column=5, value=order.get('created_at', '')[:10] if order.get('created_at') else '').border = thin_border

    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 30
    ws.column_dimensions['C'].width = 16
    ws.column_dimensions['D'].width = 14
    ws.column_dimensions['E'].width = 14

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"stats_orders_{date_from or 'all'}_{date_to or 'all'}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/products/export")
async def export_products_stats(
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    session: AsyncSession = Depends(get_session),
):
    stats = await get_all_stats(date_from, date_to, session)
    stats = stats["items"] if isinstance(stats, dict) else stats
    products_res = await session.execute(select(Product))
    products_map = {p.id: p.name for p in products_res.scalars().all()}

    wb = Workbook()
    ws = wb.active
    ws.title = "Статистика товаров"

    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill(start_color="6C63FF", end_color="6C63FF", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin'),
    )

    ws.merge_cells('A1:F1')
    ws['A1'] = f"Статистика по товарам • Период: {date_from or '—'} – {date_to or '—'}"
    ws['A1'].font = Font(bold=True, size=14)

    headers = ["Товар", "В корзину", "Заказано", "Возвраты", "Продано (период), шт", "Выручка (период), ₽"]
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=3, column=col, value=h)
        c.font = hdr_font; c.fill = hdr_fill; c.border = thin_border

    for i, item in enumerate(stats):
        row = 4 + i
        ws.cell(row=row, column=1, value=products_map.get(item['product_id'], f"ID {item['product_id']}")).border = thin_border
        ws.cell(row=row, column=2, value=item['added_to_cart']).border = thin_border
        ws.cell(row=row, column=3, value=item['ordered']).border = thin_border
        ws.cell(row=row, column=4, value=item['returned']).border = thin_border
        ws.cell(row=row, column=5, value=item.get('period_sold_qty', 0)).border = thin_border
        ws.cell(row=row, column=6, value=item.get('period_sold_sum', 0)).border = thin_border
        for col in range(2, 7):
            ws.cell(row=row, column=col).number_format = '# ##0'

    for col, width in enumerate([30, 12, 12, 12, 16, 16], 1):
        ws.column_dimensions[chr(64 + col)].width = width

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"stats_products_{date_from or 'all'}_{date_to or 'all'}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )